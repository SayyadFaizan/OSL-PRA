import os
import subprocess
import tempfile
import matplotlib.pyplot as plt
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

STUDENT_DETAILS = """Name: Sayyad Faizan
Class/Batch: TE CO - I 03
Roll No.: 24DCO11
Subject: Cloud Computing & Frameworks (CCF)
"""

def render_html_to_screenshot(html_content, output_path, width=900, height=500):
    """
    Renders HTML content to a PNG image using headless Google Chrome.
    """
    with tempfile.NamedTemporaryFile('w', suffix='.html', delete=False) as f:
        f.write(html_content)
        temp_html_path = f.name

    try:
        cmd = [
            'google-chrome',
            '--headless',
            '--disable-gpu',
            '--no-sandbox',
            f'--window-size={width},{height}',
            f'--screenshot={output_path}',
            temp_html_path
        ]
        subprocess.run(cmd, check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    finally:
        if os.path.exists(temp_html_path):
            os.remove(temp_html_path)

def create_terminal_image(title, command_text, output_text, output_path, width_px=800, height_px=400, prompt_user="ubuntu@cloud-vm:~$"):
    fig, ax = plt.subplots(figsize=(width_px/100, height_px/100), dpi=100)
    fig.patch.set_facecolor('#1e1e1e')
    ax.set_facecolor('#1e1e1e')
    ax.axis('off')

    ax.text(0.02, 0.95, f"● ● ●  Terminal - {title}", color='#a0a0a0', fontsize=11, fontweight='bold', family='monospace', transform=ax.transAxes)
    ax.plot([0, 1], [0.91, 0.91], color='#333333', transform=ax.transAxes, linewidth=1)

    full_text = f"{prompt_user} {command_text}\n" + output_text

    lines = full_text.split('\n')
    formatted_lines = []
    for line in lines:
        if len(line) > 85:
            formatted_lines.append(line[:85] + '...')
        else:
            formatted_lines.append(line)

    display_content = '\n'.join(formatted_lines[:18])
    safe_content = display_content.replace('$', '\\$')

    ax.text(0.02, 0.85, safe_content, color='#00ff66', fontsize=9.5, family='monospace', verticalalignment='top', transform=ax.transAxes)

    plt.tight_layout()
    plt.savefig(output_path, bbox_inches='tight', pad_inches=0.1, facecolor=fig.get_facecolor(), edgecolor='none')
    plt.close()

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def create_base_doc():
    doc = docx.Document()
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
    return doc

def add_cover_page(doc, exp_num, exp_title):
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_title.add_run("DEPARTMENT OF COMPUTER ENGINEERING\nCLOUD COMPUTING & FRAMEWORKS\n\n")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(14)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(0, 51, 102)

    run_exp = p_title.add_run(f"EXPERIMENT NO. {exp_num}\n")
    run_exp.font.name = "Arial"
    run_exp.font.size = Pt(18)
    run_exp.font.bold = True
    run_exp.font.color.rgb = RGBColor(180, 0, 0)

    run_main = p_title.add_run(f"{exp_title.upper()}\n\n")
    run_main.font.name = "Arial"
    run_main.font.size = Pt(16)
    run_main.font.bold = True

    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, "F2F4F7")

    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_head = p.add_run("STUDENT DETAILS\n\n")
    run_head.font.name = "Arial"
    run_head.font.size = Pt(12)
    run_head.font.bold = True
    run_head.font.color.rgb = RGBColor(0, 51, 102)

    run_info = p.add_run(
        "Name: Sayyad Faizan\n"
        "Class/Batch: TE CO - I 03\n"
        "Roll No.: 24DCO11\n"
        "Academic Year: 2024-2025\n"
        "Subject: Cloud Computing & Frameworks (CCF)\n"
    )
    run_info.font.name = "Arial"
    run_info.font.size = Pt(11)

    doc.add_page_break()

print("Headless Chrome doc_helper module initialized.")
